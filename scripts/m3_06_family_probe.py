"""M3-06 R4 / R5 产品族探针：只改打分输入文本，看 Embedding 与 Rerank 对 Q003×C005 的分数与名次变化。

组件级对照：不跑 pipeline、不改语料 / qrels / 标签，也不改任何检索参数。
候选文本仍是 `build_documents` 的同一条历史检索文本，只在其后追加一行产品族；
Query 文本保持原样（Query 侧不含族名，这一点会在输出中体现）。

复算命令： uv run python scripts/m3_06_family_probe.py
"""

from pathlib import Path

import openpyxl

from casetrace.demo import build_documents, load_validated_dataset
from casetrace.retrieval.embedding import (
    QUERY_INSTRUCTION,
    SentenceTransformerEncoder,
    rank_by_similarity,
)
from casetrace.retrieval.rerank import SentenceTransformerReranker

ROOT = Path(__file__).resolve().parents[1]
DATASET_PATH = ROOT / "data/dev/demo-v3.json"
REFERENCE_PATH = ROOT / "data/reference/封装异常_failure_modes_db_structured_v5_engineering_audited-2.xlsx"

QUERY_ID = "Q003"
CASE_ID = "C005"


def family_by_product() -> dict[str, str]:
    """产品 → 产品族名映射；主数据里有该列，但 ReferenceData 不加载，探针单独读取。"""
    workbook = openpyxl.load_workbook(REFERENCE_PATH, read_only=True, data_only=True)
    try:
        product_rows = workbook["products"].iter_rows(values_only=True)
        product_header = next(product_rows)
        product_at = product_header.index("product_id")
        family_at = product_header.index("product_family_id")
        family_id_of = {
            row[product_at]: row[family_at] for row in product_rows if row[product_at] is not None
        }

        family_rows = workbook["product_families"].iter_rows(values_only=True)
        family_header = next(family_rows)
        family_id_at = family_header.index("product_family_id")
        family_name_at = family_header.index("product_family")
        name_of = {
            row[family_id_at]: row[family_name_at]
            for row in family_rows
            if row[family_id_at] is not None
        }
    finally:
        workbook.close()
    return {product_id: name_of[family_id] for product_id, family_id in family_id_of.items()}


def rank_position(order: list[str], case_id: str) -> int:
    return order.index(case_id) + 1


def main() -> None:
    records, payload, reference = load_validated_dataset(DATASET_PATH, REFERENCE_PATH)
    documents = build_documents(records, reference)
    case_ids = sorted(documents)
    query_text = next(query["text"] for query in payload["queries"] if query["query_id"] == QUERY_ID)
    product_id = next(
        detail.product_id for detail in records["details"] if detail.case_id == CASE_ID
    )
    family_name = family_by_product()[product_id]

    family_line = f"产品族：{family_name}"
    after_texts = [
        f"{documents[case_id]}\n{family_line}" if case_id == CASE_ID else documents[case_id]
        for case_id in case_ids
    ]

    print(f"探针对象：{QUERY_ID} × {CASE_ID}（{product_id} / {family_name}）")
    print(f"Query 文本（未改）：{query_text}")
    print(f"追加行：{family_line!r}（只加给 {CASE_ID} 的候选文本）")
    print()

    encoder = SentenceTransformerEncoder()
    query_vector = encoder.encode([QUERY_INSTRUCTION + query_text])[0]
    before_vectors = encoder.encode([documents[case_id] for case_id in case_ids])
    after_vectors = encoder.encode(after_texts)
    before_similarity = before_vectors @ query_vector
    after_similarity = after_vectors @ query_vector
    before_embedding_order = [
        hit.case_id for hit in rank_by_similarity(case_ids, before_similarity, top_k=len(case_ids))
    ]
    after_embedding_order = [
        hit.case_id for hit in rank_by_similarity(case_ids, after_similarity, top_k=len(case_ids))
    ]
    print("== R4 Embedding 余弦相似度（已归一化，点积）")
    print(f"   {CASE_ID} 原文本        {before_similarity[case_ids.index(CASE_ID)]:.6f}"
          f"  名次 {rank_position(before_embedding_order, CASE_ID)}/{len(case_ids)}")
    print(f"   {CASE_ID} 追加产品族后  {after_similarity[case_ids.index(CASE_ID)]:.6f}"
          f"  名次 {rank_position(after_embedding_order, CASE_ID)}/{len(case_ids)}")
    print(f"   差值              "
          f"{after_similarity[case_ids.index(CASE_ID)] - before_similarity[case_ids.index(CASE_ID)]:+.6f}")
    print(f"   追加后前 4：{' '.join(after_embedding_order[:4])}"
          f"（原：{' '.join(before_embedding_order[:4])}）")
    print()

    reranker = SentenceTransformerReranker()
    before_rerank = reranker.score(query_text, [documents[case_id] for case_id in case_ids])
    after_rerank = reranker.score(query_text, after_texts)
    before_rerank_order = [
        hit.case_id for hit in rank_by_similarity(case_ids, before_rerank, top_k=len(case_ids))
    ]
    after_rerank_order = [
        hit.case_id for hit in rank_by_similarity(case_ids, after_rerank, top_k=len(case_ids))
    ]
    print("== R5 Rerank 分数（bge-reranker-base，Sigmoid 后）")
    print(f"   {CASE_ID} 原文本        {before_rerank[case_ids.index(CASE_ID)]:.6f}"
          f"  名次 {rank_position(before_rerank_order, CASE_ID)}/{len(case_ids)}"
          f"（M3-05 记录 0.5901 / 第 7）")
    print(f"   {CASE_ID} 追加产品族后  {after_rerank[case_ids.index(CASE_ID)]:.6f}"
          f"  名次 {rank_position(after_rerank_order, CASE_ID)}/{len(case_ids)}")
    print(f"   差值              "
          f"{after_rerank[case_ids.index(CASE_ID)] - before_rerank[case_ids.index(CASE_ID)]:+.6f}")
    print(f"   追加后前 4：{' '.join(after_rerank_order[:4])}"
          f"（原：{' '.join(before_rerank_order[:4])}）")


if __name__ == "__main__":
    main()

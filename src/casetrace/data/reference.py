"""读取演示与异常站点校验所需主数据；完整 BOM／设备导入留待实际需要。"""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class ReferenceData:
    products: dict[str, dict[str, str]]
    failure_modes: dict[str, dict[str, str]]
    product_customers: dict[str, str]
    processes: dict[str, str]
    route_processes: dict[str, set[str]]

    def validator_maps(self) -> dict:
        """从同一份主数据构造 Validator 的客户、路线、异常和工序映射。"""
        return {
            "product_customers": self.product_customers,
            "product_routes": {key: row["package_route"] for key, row in self.products.items()},
            "failure_mode_routes": {
                key: {route.strip() for route in row["applicable_package"].split(";")}
                for key, row in self.failure_modes.items()
            },
            "processes": self.processes,
            "route_processes": self.route_processes,
        }


def _read_table(workbook, name: str, key: str | tuple[str, ...], columns: tuple[str, ...]) -> dict:
    """拒绝重复单键或组合键；不把数字 ID 转成字符串来掩盖前导零丢失。"""
    if name not in workbook.sheetnames:
        raise ValueError(f"主数据缺少 Sheet：{name}")
    rows = workbook[name].iter_rows(values_only=True)
    headers = next(rows, ())
    if any(headers.count(column) != 1 for column in columns):
        raise ValueError(f"{name} 缺少必需列或列名重复：{columns}")
    positions = {column: headers.index(column) for column in columns}
    records = {}
    for number, values in enumerate(rows, start=2):
        if all(value is None for value in values):
            continue
        row = {column: values[position] for column, position in positions.items()}
        if any(not isinstance(value, str) or not value.strip() for value in row.values()):
            raise ValueError(f"{name} 第 {number} 行：必需字段须为非空文本，ID 须保留前导零")
        identity = tuple(row[column] for column in key) if isinstance(key, tuple) else row[key]
        if identity in records:
            raise ValueError(f"{name} 第 {number} 行：{key}={identity} 重复")
        records[identity] = row
    if not records:
        raise ValueError(f"主数据表 {name} 不能为空")
    return records


def load_reference(path: Path) -> ReferenceData:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        products = _read_table(workbook, "products", "product_id",
                               ("product_id", "product_name", "package_route"))
        modes = _read_table(workbook, "failure_modes", "failure_mode_id",
                           ("failure_mode_id", "failure_mode", "applicable_package",
                            "possible_root_causes", "corrective_actions"))
        assignments = _read_table(workbook, "customer_product_map", "product_id",
                                 ("product_id", "customer_id"))
        customers = _read_table(workbook, "customers", "customer_id", ("customer_id",))
        routes = _read_table(workbook, "package_routes", "package_route", ("package_route",))
        process_rows = _read_table(workbook, "process_master", "process_id", ("process_id", "process"))
        process_links = _read_table(
            workbook, "package_process_map", ("package_route", "process_id"),
            ("package_route", "process_id", "process"),
        )
    finally:
        workbook.close()

    product_customers = {key: row["customer_id"] for key, row in assignments.items()}
    if products.keys() != product_customers.keys() or set(product_customers.values()) != set(customers):
        raise ValueError("客户归属必须覆盖全部产品及客户，且引用有效")
    processes = {key: row["process"] for key, row in process_rows.items()}
    route_processes = {route: set() for route in routes}
    for (route, process_id), row in process_links.items():
        if route not in routes:
            raise ValueError(f"package_process_map 引用了未知路线：{route}")
        if process_id not in processes:
            raise ValueError(f"package_process_map 引用了未知工序：{process_id}")
        if row["process"] != processes[process_id]:
            raise ValueError(f"package_process_map 的工序名称与 process_master 不一致：{process_id}")
        route_processes[route].add(process_id)
    for route, selected in route_processes.items():
        if not selected:
            raise ValueError(f"路线 {route} 缺少工序映射")
    reference = ReferenceData(products, modes, product_customers, processes, route_processes)
    maps = reference.validator_maps()
    used_routes = set(maps["product_routes"].values())
    for applicable in maps["failure_mode_routes"].values():
        used_routes.update(applicable)
    if not used_routes <= routes.keys():
        raise ValueError("产品或 Failure Mode 引用了未知路线")
    return reference

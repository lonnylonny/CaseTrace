"""一次读取完整主数据；同一份表行用于校验、入库和重建 ReferenceData。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from casetrace.data.reference import ReferenceData
from casetrace.storage.content import INSERT_COLUMNS

SHEET_COLUMNS = {
    "customers": ("customer_id", "customer_name"),
    "product_families": ("product_family_id", "product_family"),
    "package_routes": ("package_route", "carrier", "die_interconnect",
                       "external_terminal", "description"),
    "products": ("product_id", "product_name", "product_family_id", "package_route",
                 "product_function"),
    "customer_product_map": ("customer_id", "product_id"),
    "failure_modes": ("failure_mode_id", "failure_mode", "applicable_package",
                      "applicable_process", "possible_root_causes", "failure_effects",
                      "detection_methods", "screening_methods", "corrective_actions", "notes"),
    "process_master": ("process_id", "process", "process_category", "description"),
    "package_process_map": ("package_route", "process_id", "process", "sequence_no",
                            "relation_type", "process_scope", "notes"),
}

# 摘要里主数据按自然键排序（消费方都按 key 查表，行顺序无语义）。
TABLE_KEYS = {
    "customers": ("customer_id",),
    "product_families": ("product_family_id",),
    "package_routes": ("package_route",),
    "products": ("product_id",),
    "processes": ("process_id",),
    "package_process_map": ("package_route", "process_id"),
    "failure_modes": ("failure_mode_id",),
    "failure_mode_routes": ("failure_mode_id", "package_route"),
}


def _text(value, *, sheet: str, number: int, column: str) -> str | None:
    """空白读作 None；非文本单元格明确报错，不把数字悄悄转成文本掩盖前导零丢失。"""

    if value is None:
        return None
    if isinstance(value, str):
        return value if value.strip() else None
    raise ValueError(f"{sheet} 第 {number} 行 | {column} | 必须是文本：{value!r}")


def _sequence_no(value, *, sheet: str, number: int) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{sheet} 第 {number} 行 | sequence_no | 必须是整数或空：{value!r}")
    if isinstance(value, float) and not value.is_integer():
        raise ValueError(f"{sheet} 第 {number} 行 | sequence_no | 必须是整数或空：{value!r}")
    return int(value)


def _split(value: str | None) -> list[str] | None:
    """把分号分隔的单元格拆成去空白的列表；空值返回 None。"""

    if not value:
        return None
    parts = [part.strip() for part in value.split(";") if part.strip()]
    return parts or None


def _read_sheet(workbook, name: str) -> list[dict]:
    """检查列头和单元格类型；必填、主外键由数据库约束检查。"""

    if name not in workbook.sheetnames:
        raise ValueError(f"主数据缺少 Sheet：{name}")
    columns = SHEET_COLUMNS[name]
    rows = workbook[name].iter_rows(values_only=True)
    headers = next(rows, ())
    if any(headers.count(column) != 1 for column in columns):
        raise ValueError(f"{name} 缺少必需列或列名重复：{columns}")
    positions = {column: headers.index(column) for column in columns}
    records: list[dict] = []
    for number, values in enumerate(rows, start=2):
        if all(value is None for value in values):
            continue
        row: dict = {}
        for column, position in positions.items():
            row[column] = (
                _sequence_no(values[position], sheet=name, number=number)
                if column == "sequence_no"
                else _text(values[position], sheet=name, number=number, column=column)
            )
        records.append(row)
    if not records:
        raise ValueError(f"主数据表 {name} 不能为空")
    return records


def build_reference_rows(reference_path: Path) -> tuple[dict[str, list[dict]], dict[str, list[str]]]:
    """读取一次 Excel，返回完整表行和异常适用路线的源顺序。"""

    workbook = load_workbook(reference_path, read_only=True, data_only=True)
    try:
        sheets = {name: _read_sheet(workbook, name) for name in SHEET_COLUMNS}
    finally:
        workbook.close()

    customer_by_product = {
        row["product_id"]: row["customer_id"] for row in sheets["customer_product_map"]
    }
    if len(customer_by_product) != len(sheets["customer_product_map"]):
        raise ValueError("customer_product_map 中一个产品只能对应一个客户")
    processes = {row["process_id"]: row["process"] for row in sheets["process_master"]}
    for row in sheets["package_process_map"]:
        if row["process"] != processes.get(row["process_id"]):
            raise ValueError(f"路线工序名称与 process_master 不一致：{row['process_id']}")
    products = []
    for row in sheets["products"]:
        customer_id = customer_by_product.get(row["product_id"])
        if customer_id is None:
            raise ValueError(f"customer_product_map 缺少产品 {row['product_id']} 的客户归属")
        products.append({
            "product_id": row["product_id"], "product_name": row["product_name"],
            "product_family_id": row["product_family_id"], "package_route": row["package_route"],
            "product_function": row["product_function"], "customer_id": customer_id,
        })

    failure_mode_routes = []
    route_order: dict[str, list[str]] = {}
    for row in sheets["failure_modes"]:
        routes = _split(row["applicable_package"]) or []
        route_order[row["failure_mode_id"]] = routes
        for route in routes:
            failure_mode_routes.append({
                "failure_mode_id": row["failure_mode_id"], "package_route": route,
            })

    rows = {
        table: [{column: row[column] for column in INSERT_COLUMNS[table]} for row in sheets[sheet]]
        for table, sheet in (
            ("customers", "customers"), ("product_families", "product_families"),
            ("package_routes", "package_routes"), ("processes", "process_master"),
            ("package_process_map", "package_process_map"),
        )
    }
    rows["products"] = products
    rows["failure_modes"] = [
        {column: _split(row[column]) if column == "applicable_process" else row[column]
         for column in INSERT_COLUMNS["failure_modes"]}
        for row in sheets["failure_modes"]
    ]
    rows["failure_mode_routes"] = failure_mode_routes
    for table, key in TABLE_KEYS.items():
        rows[table].sort(key=lambda row, key=key: tuple(row[column] for column in key))
    return rows, route_order


def build_reference_data(reference_rows: dict, route_order: dict[str, list[str]]) -> ReferenceData:
    """重建等价 ReferenceData；`applicable_package` 按记录的源顺序拼回。"""

    failure_modes = {}
    for row in reference_rows["failure_modes"]:
        mode_id = row["failure_mode_id"]
        routes = route_order[mode_id]
        failure_modes[mode_id] = {
            "failure_mode_id": mode_id, "failure_mode": row["failure_mode"],
            "applicable_package": "; ".join(routes),
            "possible_root_causes": row["possible_root_causes"],
            "corrective_actions": row["corrective_actions"],
        }
    route_processes = {row["package_route"]: set() for row in reference_rows["package_routes"]}
    for row in reference_rows["package_process_map"]:
        route_processes[row["package_route"]].add(row["process_id"])
    return ReferenceData(
        products={
            row["product_id"]: {
                "product_id": row["product_id"], "product_name": row["product_name"],
                "product_family_id": row["product_family_id"],
                "package_route": row["package_route"],
            }
            for row in reference_rows["products"]
        },
        failure_modes=failure_modes,
        product_customers={
            row["product_id"]: row["customer_id"] for row in reference_rows["products"]
        },
        processes={row["process_id"]: row["process"] for row in reference_rows["processes"]},
        route_processes=route_processes,
        product_families={
            row["product_family_id"]: row["product_family"]
            for row in reference_rows["product_families"]
        },
    )



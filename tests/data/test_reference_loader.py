from openpyxl import load_workbook
import pytest

from casetrace.data.reference import load_reference


def test_reads_maps_and_preserves_leading_zeros(reference_path):
    reference = load_reference(reference_path)
    assert reference.product_customers == {"P1": "CUS1"}
    assert reference.validator_maps()["failure_mode_routes"] == {"00001": {"LF_WB"}}


def test_reads_all_route_processes_without_overwriting_shared_route(reference_path):
    reference = load_reference(reference_path)
    assert reference.processes == {
        "P004": "Wire Bond", "P007": "Molding", "P013": "Storage & Transportation",
    }
    assert reference.route_processes == {"LF_WB": {"P004", "P007", "P013"}}
    assert reference.validator_maps()["processes"] == reference.processes
    assert reference.validator_maps()["route_processes"] == reference.route_processes


@pytest.mark.parametrize("sheet,row,message", [
    ("customer_product_map", ("P1", "CUS1"), "重复"),
    ("products", ("P2", "另一个产品", "LF_WB"), "覆盖全部"),
    ("products", (2, "数字 ID", "LF_WB"), "非空文本"),
    ("failure_modes", ("00002", "unknown", "UNKNOWN", "原因", "措施"), "未知路线"),
    ("process_master", ("P004", "duplicate"), "重复"),
    ("process_master", (4, "numeric ID"), "非空文本"),
    ("package_process_map", ("LF_WB", "P004", "Wire Bond"), "重复"),
    ("package_process_map", ("UNKNOWN", "P004", "Wire Bond"), "未知路线"),
    ("package_process_map", ("LF_WB", "P999", "Unknown"), "未知工序"),
    ("package_routes", ("EMPTY",), "缺少工序映射"),
])
def test_rejects_invalid_reference_rows(reference_path, sheet, row, message):
    workbook = load_workbook(reference_path)
    workbook[sheet].append(row)
    workbook.save(reference_path)
    workbook.close()
    with pytest.raises(ValueError, match=message):
        load_reference(reference_path)


@pytest.mark.parametrize("mutation,message", [
    ("name_mismatch", "工序名称"), ("missing_sheet", "缺少 Sheet"),
    ("missing_column", "缺少必需列"),
])
def test_rejects_broken_process_tables(reference_path, mutation, message):
    workbook = load_workbook(reference_path)
    if mutation == "name_mismatch":
        workbook["package_process_map"]["C2"] = "Molding"
    elif mutation == "missing_sheet":
        del workbook["process_master"]
    else:
        workbook["package_process_map"]["B1"] = "wrong_column"
    workbook.save(reference_path)
    workbook.close()
    with pytest.raises(ValueError, match=message):
        load_reference(reference_path)


def test_incoming_auxiliary_and_optional_processes_need_no_sequence(reference_path):
    workbook = load_workbook(reference_path)
    workbook["process_master"].append(("P014", "Material Income"))
    workbook["process_master"].append(("P012", "SMT"))
    workbook["package_routes"].append(("SUBSTRATE_FC",))
    links = workbook["package_process_map"]
    links["D1"] = "sequence_no"
    links["E1"] = "relation_type"
    links["E4"] = "auxiliary"
    links.append(("LF_WB", "P014", "Material Income", None, "incoming"))
    links.append(("SUBSTRATE_FC", "P012", "SMT", None, "optional"))
    workbook.save(reference_path)
    workbook.close()
    reference = load_reference(reference_path)
    assert reference.route_processes == {
        "LF_WB": {"P004", "P007", "P013", "P014"}, "SUBSTRATE_FC": {"P012"},
    }

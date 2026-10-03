"""只读检查：dev-v3 的 Query 标识是否与主数据及语料一致。

不修改任何文件；不读取 qrels 标签，只核对事实关系。
用法：uv run python tmp/check_query_enrichment.py
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATASET = ROOT / "data/dev/demo-v3.json"
REFERENCE = ROOT / "data/reference/封装异常_failure_modes_db_structured_v5_engineering_audited-2.xlsx"

PRODUCT = re.compile(r"PROD_\d+")
CUSTOMER = re.compile(r"CUS_\d+")
PRODUCTION_LOT = re.compile(r"DEV_PL_\d+")
CUSTOMER_LOT = re.compile(r"DEV_CL_\d+")


def _reference() -> tuple[dict, dict, dict]:
    workbook = openpyxl.load_workbook(REFERENCE, read_only=True, data_only=True)
    try:
        products = {row[0]: {"family": row[2], "route": row[3]}
                    for row in workbook["products"].iter_rows(min_row=2, values_only=True) if row and row[0]}
        customers = {row[0] for row in workbook["customers"].iter_rows(min_row=2, values_only=True) if row and row[0]}
        assignments = {row[1]: row[0] for row in workbook["customer_product_map"].iter_rows(min_row=2, values_only=True)
                       if row and row[1]}
        families = {row[0]: row[1] for row in workbook["product_families"].iter_rows(min_row=2, values_only=True)
                    if row and row[0]}
    finally:
        workbook.close()
    return products, customers, assignments | {"__families__": families}


def main() -> int:
    payload = json.loads(DATASET.read_text(encoding="utf-8"))
    products, customers, assignments = _reference()
    families = assignments.pop("__families__")
    details = {row["case_id"]: row for row in payload["details"]}
    lot_owner = {row["production_lot"]: row for row in payload["details"]}

    errors: list[str] = []
    print(f"语料 {DATASET.relative_to(ROOT)}  queries={len(payload['queries'])} cases={len(payload['cases'])}")
    for query in payload["queries"]:
        query_id, text = query["query_id"], query["text"]
        found = {
            "product": PRODUCT.findall(text),
            "customer": CUSTOMER.findall(text),
            "production_lot": PRODUCTION_LOT.findall(text),
            "customer_lot": CUSTOMER_LOT.findall(text),
        }
        for field, values in found.items():
            if len(values) != 1:
                errors.append(f"{query_id} | {field} | 需要恰好一个标识，实际 {values}")
        if errors and any(query_id in item for item in errors):
            continue
        product, customer = found["product"][0], found["customer"][0]
        production_lot, customer_lot = found["production_lot"][0], found["customer_lot"][0]

        if product not in products:
            errors.append(f"{query_id} | product {product} | 主数据中不存在")
            continue
        if customer not in customers:
            errors.append(f"{query_id} | customer {customer} | 主数据中不存在")
        if assignments.get(product) != customer:
            errors.append(f"{query_id} | CR-04 | {product} 的客户为 {assignments.get(product)}，与 Query 的 {customer} 不一致")

        owner = lot_owner.get(production_lot)
        if owner is None:
            reused = "新批号"
        else:
            reused = f"复用 {owner['case_id']} 的批号"
            if owner["product_id"] != product:
                errors.append(f"{query_id} | CR-08 | 生产批 {production_lot} 在语料中属 {owner['product_id']}")
            if owner["customer_lot"] != customer_lot:
                errors.append(f"{query_id} | CR-09 | 生产批 {production_lot} 在语料中的客户批为 {owner['customer_lot']}")

        same_lot_cases = sorted(case for case, row in details.items() if row["production_lot"] == production_lot)
        relations = [case for case, row in details.items()
                     if row["product_id"] == product or products[row["product_id"]]["family"] == products[product]["family"]]
        print(f"{query_id} {product}/{families[products[product]['family']]}/{products[product]['route']} "
              f"客户 {customer} | {production_lot}/{customer_lot}（{reused}）")
        print(f"    语料同批 Case: {same_lot_cases or '无'}；同产品或同产品族的 Case: {sorted(relations) or '无'}")

    print(f"\n检查结果：{'PASS，无错误' if not errors else 'FAIL'}")
    for error in errors:
        print(f"  - {error}")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())

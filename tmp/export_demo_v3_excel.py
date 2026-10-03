"""把 Development 语料 JSON 导出成 Excel 审阅视图。

只读导出：Excel 只是同一份 JSON 的表格视图，不是权威来源，编辑也不会回写。
字段含义以 docs/data/CaseTrace_Data_Structure_V2_No_Scenario.md 为准；
相关性标签 / qrels 不在本文件中，仍由用户单独确认。

用法：
    uv run python tmp/export_demo_v3_excel.py
    uv run python tmp/export_demo_v3_excel.py --data data/dev/demo-v3.json --output tmp/demo-v3.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from casetrace.data.reference import load_reference

DEFAULT_REFERENCE = Path("data/reference/封装异常_failure_modes_db_structured_v5_engineering_audited-2.xlsx")

# 语料级标量：放在“说明”工作表，不单独占列。
SCALARS = (
    ("split", "数据划分"),
    ("review_status", "语料审阅状态"),
    ("corpus_version", "语料版本"),
    ("generation_basis", "生成依据"),
    ("supersedes.path", "被取代的语料"),
    ("supersedes.sha256", "被取代语料的 SHA-256"),
)

# 每个工作表声明要导出的字段路径（点号表示嵌套字段）。
# 未声明的字段会在导出前报错，避免新增字段被静默漏掉。
COLUMNS: dict[str, tuple[str, ...]] = {
    "cases": (
        "case_id", "abnormal_description", "root_cause", "corrective_action",
        "investigation_others", "abnormal_processes",
    ),
    "details": (
        "detail_id", "case_id", "product_id", "customer_lot", "production_lot",
        "production_time", "detection_stage", "detection_time", "abnormal_types",
        "affected_qty", "disposition",
    ),
    "evidences": (
        "checkpoint_id", "case_id", "checkpoint_type", "custom_name", "result", "relevance",
    ),
    "groups": ("group_id", "group_type", "description", "other_type_description"),
    "memberships": ("group_id", "case_id", "association_reason"),
    "sources": (
        "case_id", "sheet", "failure_mode_id", "root_cause", "corrective_action",
        "closure_status", "review_status",
        "simulation_revision.date", "simulation_revision.description",
        "simulation_revision.evidence_ids",
        "generation_note.date", "generation_note.family", "generation_note.synonym_of",
        "generation_note.coverage",
    ),
    "queries": ("query_id", "known_at", "text"),
    "case_families": ("family", "cases", "note"),
}

# 推导视图：把主数据名称与 Case 事实并排，便于判断同产品 / 类似产品 / 同异常。
REFERENCE_COLUMNS: tuple[str, ...] = (
    "case_id", "detail_ids", "product_ids", "product_names", "product_families",
    "package_routes", "abnormal_type_ids", "abnormal_type_names",
    "abnormal_process_ids", "abnormal_process_names", "case_family",
)

# 字段中文说明：文字取自 docs/data 的冻结结构与 dataset_model.py，不新增业务规则。
DESCRIPTIONS: dict[str, dict[str, str]] = {
    "cases": {
        "case_id": "主键。",
        "abnormal_description": "整体异常描述。",
        "root_cause": "原因结论，可为 NDF，但必须有明确结案结论（CR-28）。",
        "corrective_action": "具体措施，可多措施（CR-29、GR-09）。",
        "investigation_others": "补充调查文本，可选。",
        "abnormal_processes": "已确认涉及异常的 process_id 无序集合，至少一个、内部不重复（CR-48～50）。",
    },
    "details": {
        "detail_id": "主键。",
        "case_id": "所属 Case。",
        "product_id": "对应 Product；客户由 Product 推导，Case 不重复保存。",
        "customer_lot": "客户批号，厂内发现也必填。",
        "production_lot": "内部生产批号。",
        "production_time": "投批日期。",
        "detection_stage": "发现阶段（IQC / In-process / OQC / Customer / Other），不是发生工序。",
        "detection_time": "异常发现／反馈日期，不得早于投批日期（CR-21）。",
        "abnormal_types": "failure_mode_id 的无序集合，至少一个、内部不重复。",
        "affected_qty": "纳入处置范围的颗数，单位固定为 ea，不是仅确认不良的数量。",
        "disposition": "最终处置文本，允许 hold / sort / rework / release 等组合。",
    },
    "evidences": {
        "checkpoint_id": "主键。",
        "case_id": "所属 Case；调查项只挂 Case，不挂 Detail。",
        "checkpoint_type": "QC / AOI / Production / OCAP / Previous/Next Lot / Monitoring / EDX / Reliability / Material / Other。",
        "custom_name": "Other 类型的具体调查名称，Other 时必填。",
        "result": "调查结果，实际生产事实写入此处。",
        "relevance": "related / not_related / uncertain；related 标签本身不能证明结果支持 Root Cause。",
    },
    "groups": {
        "group_id": "主键。",
        "group_type": "多选分组类型（CR-42）；same_abnormal_process 表示同异常站点关联（CR-51）。",
        "description": "这组 Case 的整体说明。",
        "other_type_description": "other 分组类型的解释；包含 other 时必填（CR-43）。",
    },
    "memberships": {
        "group_id": "引用 Group。",
        "case_id": "引用 Case。",
        "association_reason": "该 Case 的入组理由，非空。",
    },
    "sources": {
        "case_id": "来源记录对应的 Case（原 JSON 中的对象键）。",
        "sheet": "候选来源所在的主数据 Sheet。",
        "failure_mode_id": "选定的 Failure Mode，须属于本 Case 的 abnormal_types。",
        "root_cause": "选定的原因候选，须在主数据 possible_root_causes 中且与 Case 文本一致。",
        "corrective_action": "选定的措施候选，须在主数据 corrective_actions 中且与 Case 文本一致。",
        "closure_status": "结案状态（confirmed 时至少需要一个 related Evidence，CR-35）。",
        "review_status": "来源记录的审阅状态，当前为草稿。",
        "simulation_revision.date": "模拟事实修订日期。",
        "simulation_revision.description": "为补齐异常站点而新增的模拟事实说明，不属于旧版 Evidence。",
        "simulation_revision.evidence_ids": "被修订影响的 Evidence ID。",
        "generation_note.date": "生成说明日期。",
        "generation_note.family": "生成侧的近重复／类比家族名，供 M6 隔离使用。",
        "generation_note.synonym_of": "同义改写的措辞父项；无则为空。",
        "generation_note.coverage": "该 Case 覆盖的缺口（同义表达、技术类比、背景关联、易混淆等）。",
    },
    "queries": {
        "query_id": "主键。",
        "known_at": "该查询的时点；Query 只含此时点已知的事实，不含后来得知的当前原因。",
        "text": "查询文本。",
    },
    "case_families": {
        "family": "家族名。",
        "cases": "家族内的 Case；家族关系不自动构成检索相关性。",
        "note": "家族说明（近重复、类比或易混淆）。",
    },
    "主数据对照": {
        "case_id": "Case 主键。",
        "detail_ids": "该 Case 的 Detail；多笔 Detail 时用“；”连接。",
        "product_ids": "对应 Product；产品名、产品族、路线由主数据查询。",
        "product_names": "主数据 products.product_name。",
        "product_families": "主数据 product_families.product_family；“类似产品”按同 Product Family 判断。",
        "package_routes": "主数据 products.package_route；仅同路线不足以成立类似产品。",
        "abnormal_type_ids": "该 Case 全部 Detail 的 abnormal_types 去重集合。",
        "abnormal_type_names": "主数据 failure_modes.failure_mode 名称。",
        "abnormal_process_ids": "Case 的 abnormal_processes（已确认异常站点）。",
        "abnormal_process_names": "主数据 process_master 的工序名称。",
        "case_family": "case_families 中该 Case 所属家族；无则为空。",
    },
}


def _value(record: dict, path: str):
    """按点号路径取值；缺失或类型不符时返回 None。"""
    value = record
    for part in path.split("."):
        if not isinstance(value, dict):
            return None
        value = value.get(part)
    return value


def _cell(value):
    """列表按“；”连接；嵌套对象必须已展开成点号列，否则视为漏列。"""
    if isinstance(value, list):
        return "；".join(str(item) for item in value) if value else None
    if isinstance(value, dict):
        raise ValueError(f"字段未展开为独立列：{value}")
    return value


def _records(payload: dict, key: str) -> list[dict]:
    """实体是列表；sources 是 case_id → 记录的对象，这里补上 case_id 列。"""
    if key not in payload:
        raise ValueError(f"语料缺少顶层字段 {key}：本脚本按 dev-v3 结构导出")
    value = payload[key]
    if isinstance(value, dict):
        return [{"case_id": case_id, **row} for case_id, row in value.items()]
    if isinstance(value, list):
        return value
    raise ValueError(f"{key} 必须是列表或对象")


def _undeclared_paths(record: dict, declared: set[str], prefix: str = "") -> list[str]:
    """递归找出没有对应列的字段路径。"""
    unknown = []
    for key, value in record.items():
        path = f"{prefix}{key}"
        if isinstance(value, dict):
            unknown.extend(_undeclared_paths(value, declared, prefix=f"{path}."))
        elif path not in declared:
            unknown.append(path)
    return unknown


def _check_columns(title: str, records: list[dict], columns: tuple[str, ...]) -> None:
    declared = set(columns)
    for record in records:
        unknown = _undeclared_paths(record, declared)
        if unknown:
            raise ValueError(f"{title} 存在未导出的字段：{sorted(unknown)}")


def _product_families(reference_path: Path) -> dict[str, str]:
    """产品 → 产品族名称；产品族不在 ReferenceData 中，这里按冻结列名直接读取。"""
    workbook = load_workbook(reference_path, read_only=True, data_only=True)
    try:
        families = {
            row[0]: row[1]
            for row in workbook["product_families"].iter_rows(min_row=2, values_only=True)
            if row and row[0]
        }
        rows = workbook["products"].iter_rows(values_only=True)
        headers = next(rows)
        product_at = headers.index("product_id")
        family_at = headers.index("product_family_id")
        return {
            row[product_at]: families.get(row[family_at])
            for row in rows if row and row[product_at]
        }
    finally:
        workbook.close()


def _reference_rows(payload: dict, reference_path: Path) -> list[dict]:
    """主数据名称与 Case 事实并排；本表只是查阅用视图，不产生相关性标签。"""
    reference = load_reference(reference_path)
    product_families = _product_families(reference_path)
    family_of_case = {
        case_id: family["family"]
        for family in payload["case_families"]
        for case_id in family["cases"]
    }
    details_of_case: dict[str, list[dict]] = {}
    for detail in payload["details"]:
        details_of_case.setdefault(detail["case_id"], []).append(detail)

    def unique(values):
        return list(dict.fromkeys(value for value in values if value is not None))

    rows = []
    for case in payload["cases"]:
        case_id = case["case_id"]
        details = details_of_case.get(case_id, [])
        mode_ids = unique(mode for detail in details for mode in detail["abnormal_types"])
        process_ids = sorted(case["abnormal_processes"])
        rows.append({
            "case_id": case_id,
            "detail_ids": unique(detail["detail_id"] for detail in details),
            "product_ids": unique(detail["product_id"] for detail in details),
            "product_names": unique(reference.products[detail["product_id"]]["product_name"] for detail in details),
            "product_families": unique(product_families.get(detail["product_id"]) for detail in details),
            "package_routes": unique(reference.products[detail["product_id"]]["package_route"] for detail in details),
            "abnormal_type_ids": mode_ids,
            "abnormal_type_names": unique(reference.failure_modes[mode]["failure_mode"] for mode in mode_ids),
            "abnormal_process_ids": process_ids,
            "abnormal_process_names": [reference.processes[key] for key in process_ids],
            "case_family": family_of_case.get(case_id),
        })
    return rows


HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
HEADER_FONT = Font(bold=True)
HEADER_ALIGNMENT = Alignment(vertical="center", wrap_text=True)
BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
MIN_WIDTH = 12
MAX_WIDTH = 48


def _display_width(value) -> int:
    """中文按两个字宽估算，仅用于列宽；不影响单元格内容。"""
    text = "" if value is None else str(value)
    return sum(2 if ord(char) > 0x2E80 else 1 for char in text)


def _write_table(workbook: Workbook, title: str, columns: tuple[str, ...], rows: list[dict]):
    """写一张表：首行字段名、冻结首行、加筛选、按内容估算列宽。"""
    worksheet = workbook.create_sheet(title)
    worksheet.append(list(columns))
    for cell in worksheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
    for row in rows:
        worksheet.append([row.get(column) for column in columns])
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = BODY_ALIGNMENT
    for index, column in enumerate(columns, start=1):
        width = max([_display_width(column)] + [_display_width(row.get(column)) for row in rows])
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, MIN_WIDTH), MAX_WIDTH)
    worksheet.freeze_panes = "A2"
    if rows:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{worksheet.max_row}"
    return worksheet


def _write_overview(workbook: Workbook, *, source: str, digest: str, payload: dict,
                    row_counts: list[tuple[str, int]]) -> None:
    entries = [
        ("本文件性质", "由下列 JSON 只读导出的审阅视图；不是权威来源，编辑不回写，冲突时以 JSON 为准。"),
        ("源文件", source),
        ("源文件 SHA-256", digest),
        ("导出脚本", "tmp/export_demo_v3_excel.py"),
        ("相关性", "本文件只呈现语料事实，不含 qrels、标注理由或生成审阅信息；相关性判断仍由用户确认。"),
    ]
    entries += [(f"{path}（{label}）", _cell(_value(payload, path))) for path, label in SCALARS]

    worksheet = workbook.create_sheet("说明")
    worksheet.append(["项", "内容"])
    for item, value in entries:
        worksheet.append([item, value])
    worksheet.append([])
    worksheet.append(["工作表", "行数（不含表头）"])
    for title, count in row_counts:
        worksheet.append([title, count])

    header_rows = (1, len(entries) + 3)
    for number in header_rows:
        for cell in worksheet[number]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.row not in header_rows:
                cell.alignment = BODY_ALIGNMENT
    worksheet.column_dimensions["A"].width = 34
    worksheet.column_dimensions["B"].width = 80


def build_workbook(data_path: Path, output_path: Path, reference_path: Path) -> list[tuple[str, int]]:
    """导出并返回各工作表行数；未声明的字段在写入前直接报错。"""
    payload = json.loads(data_path.read_text(encoding="utf-8"))
    tables: dict[str, list[dict]] = {}
    for title, columns in COLUMNS.items():
        records = _records(payload, title)
        _check_columns(title, records, columns)
        tables[title] = [
            {column: _cell(_value(record, column)) for column in columns} for record in records
        ]
    reference_records = _reference_rows(payload, reference_path)
    _check_columns("主数据对照", reference_records, REFERENCE_COLUMNS)
    reference_rows = [
        {column: _cell(record.get(column)) for column in REFERENCE_COLUMNS}
        for record in reference_records
    ]

    descriptions = [
        {"工作表": title, "字段": field, "中文说明": text}
        for title, fields in DESCRIPTIONS.items()
        for field, text in fields.items()
    ]

    workbook = Workbook()
    workbook.remove(workbook.active)
    row_counts = [("字段说明", len(descriptions))]
    row_counts += [(title, len(rows)) for title, rows in tables.items()]
    row_counts.append(("主数据对照", len(reference_rows)))
    _write_overview(workbook, source=str(data_path), digest=_sha256(data_path),
                    payload=payload, row_counts=row_counts)
    _write_table(workbook, "字段说明", ("工作表", "字段", "中文说明"), descriptions)
    for title, columns in COLUMNS.items():
        _write_table(workbook, title, columns, tables[title])
    _write_table(workbook, "主数据对照", REFERENCE_COLUMNS, reference_rows)
    workbook.save(output_path)
    return row_counts


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser(description="把 Development 语料 JSON 导出成 Excel 审阅视图")
    parser.add_argument("--data", type=Path, default=Path("data/dev/demo-v3.json"))
    parser.add_argument("--output", type=Path, default=None,
                        help="默认写到与 JSON 同目录同名的 .xlsx")
    parser.add_argument("--reference", type=Path, default=DEFAULT_REFERENCE)
    args = parser.parse_args()

    output = args.output or args.data.with_suffix(".xlsx")
    row_counts = build_workbook(args.data, output, args.reference)
    print(f"已导出 {output}")
    for title, count in row_counts:
        print(f"  {title}: {count} 行")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
